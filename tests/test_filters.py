from datetime import date, timedelta
from pathlib import Path

import yaml
from openpyxl import load_workbook

from jobscraper import filters
from jobscraper.excel import write_workbook
from jobscraper.mailer import build_email
from jobscraper.models import Job
from jobscraper.sources import efinancialcareers, linkedin

ROOT = Path(__file__).resolve().parent.parent
ROLES = yaml.safe_load((ROOT / "roles.yaml").read_text())["roles"]
CFG = yaml.safe_load((ROOT / "config.yaml").read_text())
matcher = filters.RoleMatcher(ROLES)


def role_of(title):
    m = matcher.match(title)
    return m[0]["id"] if m else None


def test_role_matching():
    assert role_of("Graduate Market Risk Analyst") == "T1-1"
    assert role_of("Market Risk Analyst - Rates") == "T1-1"
    assert role_of("Quantitative Analyst, Graduate") == "T2-1"
    assert role_of("Quant Risk Analyst") in {"T2-3", "T2-1"}
    assert role_of("Model Validation Associate") == "T1-6"
    assert role_of("Graduate Product Controller") == "T1-10"
    assert role_of("P&L Analyst") == "T1-12"
    assert role_of("Operational Risk Analyst") == "T2-13"  # broad "Risk Analyst" keyword
    assert role_of("Software Engineer") is None
    assert role_of("Marketing Analyst") is None


def test_seniority():
    words = CFG["senior_title_words"]
    assert filters.is_senior("Senior Market Risk Analyst", words)
    assert filters.is_senior("Sr. Quant Analyst", words)
    assert filters.is_senior("VP, Model Validation", words)
    assert not filters.is_senior("Graduate Market Risk Analyst", words)
    assert not filters.is_senior("Leadership Programme Analyst", words)


def test_location_exclusions():
    ex = CFG["excluded_locations"]
    assert filters.is_excluded_location("Bengaluru, Karnataka, India", ex)
    assert filters.is_excluded_location("Dubai, United Arab Emirates", ex)
    assert filters.is_excluded_location("Abu Dhabi, UAE", ex)
    assert not filters.is_excluded_location("London, England, United Kingdom", ex)
    assert not filters.is_excluded_location("Indianapolis, IN", ex)
    assert not filters.is_excluded_location("Remote", ex)


def test_years_required():
    assert filters.years_required("You have 5+ years of experience in risk") == 5
    assert filters.years_required("1-2 years relevant experience") == 1
    assert filters.years_required("graduates welcome") is None


def test_sponsorship():
    assert filters.sponsorship("Visa sponsorship is available for this role.")[0] == "Yes"
    assert filters.sponsorship("We are able to sponsor Skilled Worker visas.")[0] == "Yes"
    assert filters.sponsorship("Unfortunately we are unable to offer visa sponsorship.")[0] == "No"
    assert filters.sponsorship("We do not sponsor.")[0] == "No"
    assert filters.sponsorship("You must have the right to work in the UK.")[0] == "No"
    assert filters.sponsorship("Sponsorship is not available.")[0] == "No"
    assert filters.sponsorship("Great team, hybrid working.")[0] == "Not stated"


LINKEDIN_HTML = """
<li><div class="base-card base-search-card job-search-card" data-entity-urn="urn:li:jobPosting:4012345678">
<a class="base-card__full-link" href="https://uk.linkedin.com/jobs/view/graduate-market-risk-analyst-at-barclays-4012345678?refId=x"></a>
<h3 class="base-search-card__title"> Graduate Market Risk Analyst </h3>
<h4 class="base-search-card__subtitle"><a>Barclays</a></h4>
<span class="job-search-card__location">London, England, United Kingdom</span>
<time class="job-search-card__listdate" datetime="2026-10-27">1 day ago</time>
</div></li>"""


def test_linkedin_card_parsing():
    from bs4 import BeautifulSoup
    card = BeautifulSoup(LINKEDIN_HTML, "html.parser").select("div.base-search-card, div.job-search-card")
    assert len(card) == 1
    job = linkedin._parse_card(card[0], "Market Risk Analyst")
    assert job.title == "Graduate Market Risk Analyst"
    assert job.company == "Barclays"
    assert job.url == "https://www.linkedin.com/jobs/view/4012345678"
    assert job.posted == date(2026, 10, 27)


def test_linkedin_keywords():
    assert linkedin.build_keywords(["A B", "C"]) == '"A B" OR "C"'


def test_efc_parsing():
    data = {"data": [{"id": "24710250", "title": "Model Validation Analyst", "clientBrandName": "HSBC",
                      "jobLocation": {"displayName": "London, United Kingdom"},
                      "postedDate": "2026-10-26T09:00:00Z", "detailsPageUrl": "/jobs-UK-London-Model_Validation.id24710250",
                      "summary": "<p>We offer visa sponsorship</p>"}]}
    raws = efinancialcareers._find_job_dicts(data)
    job = efinancialcareers._to_job(raws[0], ROLES[5])
    assert job.company == "HSBC"
    assert job.location == "London, United Kingdom"
    assert job.posted == date(2026, 10, 26)
    assert job.url.startswith("https://www.efinancialcareers.com/jobs-UK")
    assert "visa sponsorship" in job.description


def test_excel_and_email(tmp_path):
    jobs = [Job("LinkedIn", "Market Risk Analyst", "Barclays", "London", "https://x/1",
                posted=date.today() - timedelta(days=1), role_title="Market Risk Analyst", tier=1,
                category="Market Risk", sponsorship="Yes", sources=["LinkedIn"])]
    stats = {"LinkedIn": {"found": 1, "errors": 0}}
    path = write_workbook(jobs, {jobs[0].key}, stats, tmp_path / "t.xlsx")
    ws = load_workbook(path)["Jobs"]
    assert ws["A2"].value == "NEW" and ws["E2"].value == "Market Risk Analyst"
    assert ws["O2"].hyperlink.target == "https://x/1"
    msg = build_email(jobs, {jobs[0].key}, stats, path, "a@gmail.com", "b@gmail.com")
    assert msg["To"] == "b@gmail.com"
    assert "1 new" in msg["Subject"]


def test_target_firms():
    from jobscraper import institutions
    assert institutions.match("Barclays") == "Barclays"
    assert institutions.match("Goldman Sachs Asset Management") == "Goldman Sachs"
    assert institutions.match("Two Sigma Investments, LP") == "Two Sigma"
    assert institutions.match("Jane Street Capital") == "Jane Street"
    assert institutions.match("Acme Widgets Ltd") == ""


def test_rotation():
    from jobscraper.main import rotate
    items = list(range(10))
    seen = set()
    for slot in range(4):
        seen |= set(rotate(items, 3, slot))
    assert seen == set(items)
    assert rotate(items, 20, 5) == items


def test_one_word_keywords_need_the_phrase():
    assert role_of("Associate, Fundraise Management, Liquid Credit Strategies") is None
    assert role_of("Trading Systems Engineer - Global Options") is None
    assert role_of("Credit Analyst - Leveraged Finance") == "T1-15"
    assert role_of("Graduate Treasury Analyst") == "T2-9"


def test_rules_rechecked_on_saved_jobs():
    from jobscraper.main import rule_broken
    regs = fake_registers()
    ok = Job("LinkedIn", "Graduate Market Risk Analyst", "Barclays", "London, UK", "u",
             posted=date.today(), first_seen=date.today().isoformat())
    ok.employer_sponsor = regs.status(ok.company, ok.location)
    assert rule_broken(ok, CFG, matcher, regs) is None
    for change, reason in [
        (dict(title="Senior Market Risk Analyst"), "too senior"),
        (dict(title="Market Risk Analyst II"), "too senior"),
        (dict(location="Mumbai, India"), "excluded country"),
        (dict(years_required=5), "asks for 5+ years"),
        (dict(sponsorship="No"), "no sponsorship"),
        (dict(posted=date.today() - timedelta(days=9)), "older than the age limit"),
        (dict(title="Software Engineer"), "title doesn't match a target role"),
    ]:
        job = Job(**{**ok.__dict__, **change})
        assert rule_broken(job, CFG, matcher, regs) == reason, change


def test_recruiters_and_duplicates():
    assert filters.is_recruiter("Oxford Knight")
    assert not filters.is_recruiter("Goldman Sachs")
    a = Job("LinkedIn", "Market Risk Analyst", "JPMorgan Chase", "", "u")
    b = Job("Careers", "Market Risk Analyst", "JPMorganChase", "", "v")
    assert a.key == b.key


def fake_registers():
    from jobscraper.sponsor_registers import SponsorRegisters, clean
    regs = SponsorRegisters.__new__(SponsorRegisters)
    regs.lists = {"UK": {clean("Barclays Bank PLC")}, "US": {clean("JPMORGAN CHASE BANK, N.A."), clean("Goldman Sachs & Co. LLC")}, "NL": set()}
    regs.loaded = {"UK": True, "US": True, "NL": False}
    return regs


def test_sponsorship_filter():
    from jobscraper.main import sponsorship_rule
    regs = fake_registers()

    def check(company, location, sponsorship="Not stated", mode="posting_or_employer"):
        job = Job("LinkedIn", "Market Risk Analyst", company, location, "u", sponsorship=sponsorship)
        job.employer_sponsor = regs.status(company, location)
        return sponsorship_rule(job, mode, regs)

    assert check("Barclays", "London, England, United Kingdom") is None          # UK list, UK job
    assert check("JPMorganChase", "New York, NY") is None                        # US list, US job
    assert check("Goldman Sachs", "Jersey City, NJ") is None                     # "Goldman Sachs & Co. LLC"
    assert check("Barclays", "New York, NY") == "employer not on a sponsor list"  # on UK list only, US job
    assert check("Barclays", "Singapore") is None                                # no list for Singapore
    assert check("Tiny Fintech Ltd", "London, UK") == "employer not on a sponsor list"
    assert check("Tiny Fintech Ltd", "London, UK", sponsorship="Yes") is None     # posting offers it
    assert check("Barclays", "London, UK", mode="posting_only") == "posting doesn't offer sponsorship"
    assert check("Tiny Fintech Ltd", "London, UK", mode="lenient") is None

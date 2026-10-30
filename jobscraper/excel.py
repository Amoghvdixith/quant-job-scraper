from __future__ import annotations

from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Job

COLUMNS = [
    ("New", 7), ("Tier", 6), ("Category", 22), ("Matched Role", 30), ("Job Title", 45),
    ("Company", 28), ("On Your Target List", 22), ("Location", 30), ("Posted", 11), ("Days Old", 9), ("Source", 18),
    ("Visa Sponsorship", 16), ("Employer On Sponsor List", 22), ("Years Exp. Asked", 15), ("Qualification Asked", 16),
    ("Apply Link", 60), ("Apply On", 18), ("Job Listing", 12), ("Posted By", 18), ("Description Checked", 12), ("Sponsorship Evidence", 60),
]
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
NEW_FILL = PatternFill("solid", fgColor="E2EFDA")
YES_FILL = PatternFill("solid", fgColor="C6EFCE")


def write_workbook(jobs: list[Job], new_keys: set[str], stats: dict, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append([c for c, _ in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    today = date.today()
    for job in jobs:
        is_new = job.key in new_keys
        ws.append([
            "NEW" if is_new else "",
            job.tier,
            job.category,
            job.role_title,
            job.title,
            job.company,
            job.target_firm,
            job.location,
            job.posted,
            (today - job.posted).days if job.posted else None,
            ", ".join(job.sources or [job.source]),
            job.sponsorship,
            job.employer_sponsor or "Not found",
            job.years_required,
            job.qualification or "",
            job.apply_link,
            job.apply_on,
            "View",
            "Recruitment agency" if job.recruiter else "Employer",
            "Yes" if job.description_checked else "No",
            job.sponsorship_evidence,
        ])
        row = ws.max_row
        link = ws.cell(row=row, column=16)
        link.hyperlink = job.apply_link
        link.font = Font(color="0563C1", underline="single")
        listing = ws.cell(row=row, column=18)
        listing.hyperlink = job.url
        listing.font = Font(color="0563C1", underline="single")
        ws.cell(row=row, column=9).number_format = "yyyy-mm-dd"
        if is_new:
            ws.cell(row=row, column=1).fill = NEW_FILL
        if job.sponsorship == "Yes":
            ws.cell(row=row, column=12).fill = YES_FILL
        if job.employer_sponsor:
            ws.cell(row=row, column=13).fill = YES_FILL
        if job.target_firm:
            ws.cell(row=row, column=7).fill = YES_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Run Summary")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 18
    rows = [
        ("Run time (UTC)", datetime.now(timezone.utc).replace(tzinfo=None).strftime("%Y-%m-%d %H:%M")),
        ("Jobs in this sheet", len(jobs)),
        ("New since last email", len(new_keys)),
        ("Offer visa sponsorship", sum(j.sponsorship == "Yes" for j in jobs)),
        ("At your target institutions", sum(bool(j.target_firm) for j in jobs)),
        ("Posted by recruitment agencies", sum(j.recruiter for j in jobs)),
        ("", ""),
        ("Source", "Jobs found this run"),
    ]
    rows += [(name, s["found"]) for name, s in stats.items()]
    rows += [("", ""), ("Source", "Failed requests")]
    rows += [(name, s["errors"]) for name, s in stats.items()]
    for r in rows:
        summary.append(list(r))
        if r[0] == "Source":
            for cell in summary[summary.max_row]:
                cell.font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
